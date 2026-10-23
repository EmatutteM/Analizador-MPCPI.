import streamlit as st
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

st.set_page_config(page_title="Análisis de Precios de Transferencia", page_icon="📊", layout="centered")

st.title("📊 Análisis de Precios de Transferencia (PCPI / CUP)")
st.write("Sube la plantilla de Excel con las operaciones para generar el informe comparativo formateado.")

archivo_subido = st.file_uploader("Cargar archivo de operaciones (.xlsx)", type=["xlsx"])

def procesar_excel(file_bytes):
    xls = pd.ExcelFile(file_bytes)
    pestaña_origen = 'operaciones' if 'operaciones' in xls.sheet_names else xls.sheet_names[0]
    df = pd.read_excel(file_bytes, sheet_name=pestaña_origen)
    
    df.columns = df.columns.astype(str).str.strip()
    
    mapeo = {
        'Artículo': 'descripcion_producto',
        'articulo': 'descripcion_producto',
        'Precio Unitario USD FOB': 'precio_fob_usd',
        'precio unitario usd fob': 'precio_fob_usd',
        'Vinculación': 'vinculacion',
        'vinculacion': 'vinculacion',
        'Fecha': 'fecha',
        'Cliente': 'cliente',
        'País': 'pais',
        'Cantidad': 'cantidad'
    }
    df = df.rename(columns=mapeo)
    
    if 'fecha' in df.columns:
        df['fecha'] = pd.to_datetime(df['fecha']).dt.strftime('%Y-%m-%d')
        
    df['vinc_clean'] = df['vinculacion'].astype(str).str.strip().str.lower()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Análisis de Precios"
    
    blue_fill = PatternFill(start_color="337AB7", end_color="337AB7", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    green_font = Font(name="Calibri", size=11, bold=True, color="2E7D32")
    red_font = Font(name="Calibri", size=11, bold=True, color="C62828")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000')
    )
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    headers = ["Descripcion de producto", "vinculacion", "Cliente", "País", "Fecha", "Cantidad", "Precio Unitario USD FOB", "Rango Intercuartil / Precio Comparable", ""]
    ws.append(headers)
    ws.merge_cells("H1:I1")
    ws["H1"] = "Rango Intercuartil / Precio Comparable"
    
    for col_num in range(1, 10):
        cell = ws.cell(row=1, column=col_num)
        cell.fill, cell.font, cell.alignment, cell.border = blue_fill, white_bold, center_align, thin_border
        
    current_row = 2
    productos = df['descripcion_producto'].dropna().unique()
    
    for prod in productos:
        df_prod = df[df['descripcion_producto'] == prod]
        df_ind = df_prod[df_prod['vinc_clean'].str.contains('tercero|independiente', regex=True, na=False)]
        df_vinc = df_prod[df_prod['vinc_clean'].str.contains('relacionada|vinculada', regex=True, na=False)]
        
        if df_vinc.empty:
            continue
            
        start_row_prod = current_row
        precios_terceros = []
        
        for _, r in df_ind.iterrows():
            ws.cell(row=current_row, column=2, value="Tercero Independiente").alignment = left_align
            ws.cell(row=current_row, column=3, value=r['cliente']).alignment = left_align
            ws.cell(row=current_row, column=4, value=r['pais']).alignment = center_align
            ws.cell(row=current_row, column=5, value=r['fecha']).alignment = center_align
            ws.cell(row=current_row, column=6, value=r['cantidad']).alignment = right_align
            p_val = float(r['precio_fob_usd'])
            precios_terceros.append(p_val)
            ws.cell(row=current_row, column=7, value=p_val).number_format = '$#,##0.00'
            current_row += 1
            
        n_ind = len(precios_terceros)
        
        if n_ind == 1:
            precio_comp_unico = precios_terceros[0]
            ws.cell(row=start_row_prod, column=8, value="Precio Comparable").font = bold_font
            ws.cell(row=start_row_prod, column=9, value=precio_comp_unico).number_format = '$#,##0.00'
            if current_row < start_row_prod + 1: current_row = start_row_prod + 1
        elif n_ind > 1:
            q1, mediana, q3 = np.percentile(precios_terceros, 25), np.median(precios_terceros), np.percentile(precios_terceros, 75)
            ws.cell(row=start_row_prod, column=8, value="Primer cuartil").font = regular_font
            ws.cell(row=start_row_prod, column=9, value=q1).number_format = '$#,##0.00'
            ws.cell(row=start_row_prod+1, column=8, value="Mediana").font = regular_font
            ws.cell(row=start_row_prod+1, column=9, value=mediana).number_format = '$#,##0.00'
            ws.cell(row=start_row_prod+2, column=8, value="Tercer cuartil").font = regular_font
            ws.cell(row=start_row_prod+2, column=9, value=q3).number_format = '$#,##0.00'
            if current_row < start_row_prod + 3: current_row = start_row_prod + 3
        else:
            ws.cell(row=start_row_prod, column=8, value="Precio Comparable").font = regular_font
            ws.cell(row=start_row_prod, column=9, value="Sin comparables").alignment = center_align
            current_row = start_row_prod + 1
            
        ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=9)
        res_header = ws.cell(row=current_row, column=8, value="Resultado")
        res_header.fill, res_header.font, res_header.alignment = blue_fill, white_bold, center_align
        ws.cell(row=current_row, column=9).border = thin_border
        current_row += 1
        
        for _, r in df_vinc.iterrows():
            ws.cell(row=current_row, column=2, value="Relacionada").alignment = left_align
            ws.cell(row=current_row, column=3, value=r['cliente']).alignment = left_align
            ws.cell(row=current_row, column=4, value=r['pais']).alignment = center_align
            ws.cell(row=current_row, column=5, value=r['fecha']).alignment = center_align
            ws.cell(row=current_row, column=6, value=r['cantidad']).alignment = right_align
            p_vinc = float(r['precio_fob_usd'])
            ws.cell(row=current_row, column=7, value=p_vinc).number_format = '$#,##0.00'
            
            if n_ind == 1:
                cumple = (p_vinc >= precio_comp_unico)
            elif n_ind > 1:
                cumple = (p_vinc >= q1)
            else:
                cumple = None
                
            veredicto = "OK" if cumple else ("MAL" if cumple is False else "SIN COMPARABLES")
            
            ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=9)
            v_cell = ws.cell(row=current_row, column=8, value=veredicto)
            v_cell.alignment = center_align
            v_cell.font = green_font if veredicto == "OK" else (red_font if veredicto == "MAL" else bold_font)
            current_row += 1
            
        end_row_prod = current_row - 1
        ws.merge_cells(start_row=start_row_prod, start_column=1, end_row=end_row_prod, end_column=1)
        prod_cell = ws.cell(row=start_row_prod, column=1, value=prod)
        prod_cell.alignment, prod_cell.font = center_align, bold_font
        
        for r_i in range(start_row_prod, current_row):
            for c_i in range(1, 10): ws.cell(row=r_i, column=c_i).border = thin_border
        current_row += 1

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 18
    ws.column_dimensions['A'].width = 32
    
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

if archivo_subido is not None:
    if st.button("🚀 Generar Análisis Económico"):
        with st.spinner("Procesando operaciones y formateando reporte..."):
            excel_resultado = procesar_excel(archivo_subido)
            st.success("¡Análisis completado con éxito!")
            st.download_button(
                label="📥 Descargar Excel de Resultados",
                data=excel_resultado,
                file_name="Resultado_Precios_Transferencia.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )